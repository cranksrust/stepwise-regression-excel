"""Builds docs/sample/home-prices.xlsx: fictional home sale data for trying the add-in.

The data are simulated (seed 0). PRICE depends on SQFT, AGE, MILES_DOWNTOWN and
SCHOOL_RATING. TAX_ASSESSED is a noisy proxy built from the same drivers, so it
enters first and is later removed by stepwise. BEDROOMS, BATHROOMS and GARAGE
are correlated with size or unrelated noise.

Usage: python make_sample.py output.xlsx
"""
import sys

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAMES = ["SQFT", "BEDROOMS", "BATHROOMS", "AGE", "GARAGE", "MILES_DOWNTOWN", "SCHOOL_RATING", "TAX_ASSESSED"]


def generate(seed=0, n=60):
    rng = np.random.default_rng(seed)
    sqft = rng.normal(1900, 450, n).clip(850, 3600).round(-1)
    beds = np.clip(np.round(sqft / 600 + rng.normal(0, 0.6, n)), 1, 6)
    baths = np.clip(np.round((sqft / 900 + rng.normal(0, 0.45, n)) * 2) / 2, 1, 4.5)
    age = rng.integers(0, 70, n).astype(float)
    garage = rng.integers(0, 4, n).astype(float)
    dist = np.round(rng.uniform(2, 30, n), 1)
    school = np.round(rng.uniform(4, 10, n), 1)
    core = 0.12 * sqft + 11 * school - 3.5 * dist
    tax = np.round(0.9 * core + rng.normal(0, 18, n) + 40, 1)
    price = (70 + core - 0.9 * age + rng.normal(0, 22, n)).round(1)
    return price, np.column_stack([sqft, beds, baths, age, garage, dist, school, tax])


def build(path):
    price, X = generate()
    wb = Workbook()
    orange, tint = "BF5700", "F8E5D6"
    base = Font(name="Calibri", size=11)

    ws = wb.active
    ws.title = "Home Prices"
    header = ["PRICE"] + NAMES
    for c, h in enumerate(header, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor=tint)
        cell.border = Border(bottom=Side(style="medium", color=orange))
        cell.alignment = Alignment(horizontal="center")
    fmts = ["#,##0.0", "#,##0", "0", "0.0", "0", "0", "0.0", "0.0", "#,##0.0"]
    for r in range(len(price)):
        row = [float(price[r])] + [float(v) for v in X[r]]
        for c, v in enumerate(row, 1):
            cell = ws.cell(r + 2, c, int(v) if fmts[c - 1] in ("0", "#,##0") else v)
            cell.number_format = fmts[c - 1]
            cell.font = base
    for c in range(1, len(header) + 1):
        ws.column_dimensions[chr(64 + c)].width = 17
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Read Me", 0)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 80
    title = info.cell(1, 1, "Home Prices: sample data for the Stepwise Regression add-in")
    title.font = Font(name="Calibri", size=14, bold=True, color=orange)
    info.cell(2, 1, "Fictional data: 60 simulated home sales, made up for practice.").font = Font(name="Calibri", size=11, italic=True)

    rows = [
        ("Variable", "Meaning"),
        ("PRICE", "Sale price, $000 (the Y variable)"),
        ("SQFT", "Finished square feet"),
        ("BEDROOMS", "Number of bedrooms"),
        ("BATHROOMS", "Number of bathrooms (half baths count 0.5)"),
        ("AGE", "Age of the house in years"),
        ("GARAGE", "Garage spaces"),
        ("MILES_DOWNTOWN", "Driving distance to downtown, miles"),
        ("SCHOOL_RATING", "Rating of the assigned elementary school, 1 to 10"),
        ("TAX_ASSESSED", "County tax-assessed value, $000"),
        ("", ""),
        ("Try it", ""),
        ("1.", "Open the Home Prices sheet and click Stepwise Regression on the Home tab."),
        ("2.", "Y range: select A1:A61 and click Use selection."),
        ("3.", "X range: select B1:I61 and click Use selection."),
        ("4.", "Leave the defaults (Stepwise, T-Stat, 2.00 and 2.00) and click Run regression."),
        ("5.", "Run it again with Forward selection, then Backward elimination, and compare the Selection steps."),
        ("", ""),
        ("What to look for", ""),
        ("", "TAX_ASSESSED enters first because it tracks price closely. Once SQFT, MILES_DOWNTOWN and the other"),
        ("", "real drivers are in the model it adds nothing new, and stepwise removes it. Forward selection never"),
        ("", "re-checks, so it keeps TAX_ASSESSED, while backward elimination drops it early. Same data, three"),
        ("", "paths, and the final models are not all the same."),
    ]
    for i, (a, b) in enumerate(rows, 4):
        ca, cb = info.cell(i, 1, a), info.cell(i, 2, b)
        ca.font = cb.font = base
        if a in ("Variable", "Try it", "What to look for"):
            ca.font = Font(name="Calibri", size=11, bold=True, color=orange if a != "Variable" else "000000")
            if a == "Variable":
                cb.font = Font(name="Calibri", size=11, bold=True)
                for cell in (ca, cb):
                    cell.fill = PatternFill("solid", fgColor=tint)
                    cell.border = Border(bottom=Side(style="medium", color=orange))
    info.sheet_view.showGridLines = False
    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    build(sys.argv[1])
